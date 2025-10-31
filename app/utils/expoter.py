import logging

from docx import Document
from pathlib import Path

logger = logging.getLogger(__name__)


def generate_resume_docx(
    data: dict, lang: str = "en", title_file: str | None = None
) -> Path:
    """Create a DOCX resume file from structured resume data."""
    logger.debug("Starting DOCX generation (lang=%s)", lang)
    doc = Document()

    doc.add_heading(data.get("summary", "Resume"), level=1)
    doc.add_paragraph(" ")

    for section in data["resume_text"].split("## "):
        section = section.strip()
        if not section:
            continue
        lines = section.split("\n")
        title = lines[0].strip()
        doc.add_heading(title, level=2)
        for line in lines[1:]:
            if line.startswith("- "):
                p = doc.add_paragraph(line[2:], style="List Bullet")
            else:
                p = doc.add_paragraph(line)

    name = (
        f"resume_agent_demo_{title_file.replace(' ', '_')}.docx"
        if title_file
        else "resume_agent_demo.docx"
    )
    path = Path(f"files/{name}")
    doc.save(path)
    logger.info("DOCX resume saved to %s", path)
    return path
